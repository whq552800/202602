"""Build the reordered supplementary appendix and S1-S22 CSV set.

The appendix follows the revised main-text logic:
- data/provenance first,
- typology and temporal pathway evidence second,
- residual-fingerprint evidence third,
- sensitivity and age-boundary analyses last.

The script uses only package-local inputs and writes new outputs; it does not
modify the author-provided source Word files.
"""

from __future__ import annotations

import math
import os
import re
import zipfile
from pathlib import Path

import pandas as pd
from config import DATA_ANALYSIS, DATA_INTERMEDIATE, DATA_TABLES, PACKAGE_ROOT
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


DEFAULT_OUTPUT_DOCX = PACKAGE_ROOT / "supplement" / "LRI_supplement_v6_submit_ready.docx"
OUTPUT_TABLE_DIR = PACKAGE_ROOT / "data" / "supplement_reordered"
SUPP_FIG_DIR = PACKAGE_ROOT / "figures" / "supplement"

EN_FONT = "Times New Roman"
ZH_FONT = "SimSun"
BLACK = RGBColor(0, 0, 0)


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, encoding="utf-8-sig")


def write_csv(df: pd.DataFrame, filename: str) -> Path:
    OUTPUT_TABLE_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_TABLE_DIR / filename
    df.to_csv(path, index=False, encoding="utf-8-sig")
    return path


def clean_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    text = str(value)
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace("HAP-related", "household-energy-profile")
    text = text.replace("HAP profile", "household-energy profile")
    text = text.replace("source file:", "field/source detail:")
    text = text.replace("local shapefile:", "boundary layer:")
    text = text.replace("local file:", "crosswalk:")
    text = text.replace("combined locally", "combined for this analysis")
    text = text.replace("Figure 4", "Figure 3")
    text = text.replace("Removed from Figure 3 and current profile interpretation", "Removed from contextual residual profiling and the current residual-fingerprint figure")
    text = re.sub(r"\b[A-Z]:\\[^;,\n]+", "[local path removed]", text)
    return text


def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if out[col].dtype == object:
            out[col] = out[col].map(clean_text)
    return out


def make_s1_data_sources() -> pd.DataFrame:
    df = read_csv(DATA_TABLES / "supplement_table_s18_data_source_provenance_clean_phrasing_20260515.csv")
    df = clean_dataframe(df)
    df["Indicator code or field"] = (
        df["Indicator code or local field"]
        .astype(str)
        .str.replace(r";?\s*field/source detail:[^;]+", "", regex=True)
        .str.replace(r";?\s*processed file:[^;]+", "", regex=True)
        .str.replace(r";?\s*boundary layer:[^;]+", "", regex=True)
        .str.replace(r";?\s*crosswalk:[^;]+", "", regex=True)
        .str.replace(r"\s*;\s*$", "", regex=True)
        .str.strip()
    )
    df["Processing in this study"] = (
        df["Processing in this study"]
        .astype(str)
        .str.replace("Downloaded in four GBD Results Tool batches, combined for this analysis, and filtered", "Downloaded from the GBD Results Tool and filtered", regex=False)
        .str.replace("locally assembled", "assembled", regex=False)
        .str.replace("Figure 4", "Figure 3", regex=False)
        .str.replace("Removed from Figure 3 and current profile interpretation because it is an outcome-side measure rather than background context.", "Removed from contextual residual profiling and the current residual-fingerprint figure because it is an outcome-side measure rather than background context.", regex=False)
    )
    keep = [
        "Data domain",
        "Variable(s) in manuscript",
        "Source and dataset",
        "Indicator code or field",
        "Years used",
        "Processing in this study",
        "Manuscript-style source citation",
    ]
    return df[keep]


def make_s3_typology_summary() -> pd.DataFrame:
    df = read_csv(DATA_TABLES / "table1_primary_typology_summary_2023_no_hap_screen_20260515.csv")
    return clean_dataframe(df)


def make_s10_country_fingerprints() -> pd.DataFrame:
    residual = read_csv(DATA_TABLES / "supplement_table_s16_high_excess_contextual_residuals_no_hap_screen_20260515.csv")
    stability = read_csv(DATA_TABLES / "supplement_table_s13_primary_typology_membership_stability.csv")
    stability = stability.rename(columns={"ISO3": "iso3", "Location": "location_name"})
    stability["Primary typology"] = (
        stability["Primary typology"]
        .astype(str)
        .str.replace(", HAP-related", "", regex=False)
        .str.replace(", other", "", regex=False)
    )
    keep = [
        "iso3",
        "Primary typology",
        "High-excess stability (%)",
        "Observed rate",
        "Expected rate",
        "Excess rate",
        "Observed/expected ratio",
        "SDI",
        "HAP-PM",
        "Ambient PM2.5",
        "Clean cooking (%)",
        "PCV3 (%)",
        "Hib3 (%)",
        "DPT (%)",
        "UHC index",
        "TB incidence per 100,000",
    ]
    merged = residual.merge(stability[keep], on="iso3", how="left")
    merged["Selection frequency across 84 scenarios (%)"] = (
        merged["stable_high_excess_probability"] * 100
    ).round(1)
    columns = [
        "iso3",
        "location_name",
        "Primary typology",
        "Selection frequency across 84 scenarios (%)",
        "Observed rate",
        "Expected rate",
        "Excess rate",
        "Observed/expected ratio",
        "SDI",
        "HAP-PM",
        "Ambient PM2.5",
        "Clean cooking (%)",
        "PCV3 (%)",
        "Hib3 (%)",
        "DPT (%)",
        "UHC index",
        "TB incidence per 100,000",
        "max_contextual_domain",
        "max_contextual_signature",
        "max_contextual_z",
        "n_contextual_domains_z_ge_1",
        "contextual_domains_z_ge_1",
        "top_contextual_signatures_z_ge_1",
    ]
    return clean_dataframe(merged[columns])


TABLE_SPECS = [
    ("S1", "Data sources, variable definitions, source fields, and processing rules", make_s1_data_sources, None),
    ("S2", "Contextual-indicator coverage and source-year recency", lambda: read_csv(DATA_TABLES / "supplement_table_s12_context_indicator_coverage.csv"), None),
    ("S3", "Full under-5 typology group summary, 2023", make_s3_typology_summary, None),
    ("S4", "Structural-disadvantage flags among high-observed/not-high-excess locations", lambda: read_csv(DATA_TABLES / "supplement_table_s8_structural_disadvantage_flags_2023.csv"), None),
    ("S5", "Annual observed/excess typology counts and group summaries, 1990-2023", lambda: read_csv(DATA_TABLES / "supplement_table_s19_temporal_typology_annual_counts.csv"), None),
    ("S6", "Country-year metrics for the 2000 high-observed/high-excess cohort", lambda: read_csv(DATA_INTERMEDIATE / "cohort_2000_double_high_country_year_metrics.csv"), 18),
    ("S7", "SDI-upshift residual-excess pathway from 2000 double-high to 2023 high-excess/not-high-observed", lambda: read_csv(DATA_INTERMEDIATE / "sdi_upshift_residual_excess_pathway.csv"), None),
    ("S8", "Pathway-specific changes in adverse contextual-domain scores, 2000-2023", lambda: read_csv(DATA_INTERMEDIATE / "pathway_sdi_profile_change_2000_2023.csv"), None),
    ("S9", "Summary of adverse SDI-year-adjusted contextual residual domains among 2023 high-observed/high-excess locations", lambda: read_csv(DATA_TABLES / "supplement_table_s17_high_excess_screen_summary.csv"), None),
    ("S10", "Country-level SDI-year-adjusted residual fingerprints and selection stability among 2023 high-observed/high-excess locations", make_s10_country_fingerprints, 12),
    ("S11", "Main-model threshold and excess-rule sensitivity", lambda: read_csv(DATA_TABLES / "supplement_table_s1_threshold_sensitivity.csv"), None),
    ("S12", "Stability of clean-cooking and HAP-PM profile contrasts across model variants", lambda: read_csv(DATA_TABLES / "supplement_table_s2_profile_contrast_stability.csv"), None),
    ("S13", "Expected-burden model validation and calibration metrics", lambda: read_csv(DATA_TABLES / "supplement_table_s10_expected_model_validation.csv"), None),
    ("S14", "Population-weighted WLS sensitivity overlap", lambda: read_csv(DATA_TABLES / "supplement_table_s11_population_weighted_sensitivity.csv"), None),
    ("S15", "Age-boundary observed/excess typology summary", lambda: read_csv(DATA_TABLES / "supplement_table_s3_age_boundary_summary.csv"), None),
    ("S16", "Cross-age overlap of under-5 high-observed/high-excess locations", lambda: read_csv(DATA_TABLES / "supplement_table_s4_under5_cross_age_overlap.csv"), None),
    ("S17", "Granular age-boundary typology summary", lambda: read_csv(DATA_TABLES / "supplement_table_s6_granular_age_boundary_summary.csv"), None),
    ("S18", "Granular cross-age overlap of high-observed/high-excess membership", lambda: read_csv(DATA_TABLES / "supplement_table_s7_granular_age_overlap.csv"), 18),
    ("S19", "Duan smearing-corrected expected-rate sensitivity", lambda: read_csv(OUTPUT_TABLE_DIR / "table_s19_duan_smearing_corrected_expected_rate_sensitivity.csv"), None),
    ("S20", "2000-anchored fixed-threshold sensitivity", lambda: read_csv(OUTPUT_TABLE_DIR / "table_s20_2000_anchored_fixed_threshold_sensitivity.csv"), None),
    ("S21", "Small-population exclusion sensitivity", lambda: read_csv(OUTPUT_TABLE_DIR / "table_s21_small_population_exclusion_sensitivity.csv"), None),
    ("S22", "Harmonised 2018-2023 contextual-window sensitivity", lambda: read_csv(OUTPUT_TABLE_DIR / "table_s22_harmonised_2018_2023_contextual_window_sensitivity.csv"), None),
]


def set_run_font(run, size: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = EN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ZH_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if bold is not None:
        run.bold = bold


def apply_paragraph_font(paragraph, size: float = 10.5, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def set_doc_style(doc: Document) -> None:
    for style in doc.styles:
        if hasattr(style, "font"):
            style.font.name = EN_FONT
            style.font.color.rgb = BLACK
            style._element.rPr.rFonts.set(qn("w:eastAsia"), ZH_FONT)
    doc.styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.bold = True
        style.font.color.rgb = BLACK
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def add_para(doc: Document, text: str, style: str | None = None, size: float = 10.5, bold: bool | None = None):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def format_value(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float):
        if abs(value) >= 100:
            return f"{value:.1f}"
        if abs(value) >= 10:
            return f"{value:.2f}"
        return f"{value:.3g}"
    return clean_text(value)


def add_table_preview(doc: Document, df: pd.DataFrame, preview_rows: int | None) -> None:
    preview = df.head(preview_rows) if preview_rows else df
    if preview.empty:
        add_para(doc, "No rows available.")
        return
    table = doc.add_table(rows=1, cols=len(preview.columns))
    table.style = "Table Grid"
    for idx, column in enumerate(preview.columns):
        cell = table.rows[0].cells[idx]
        cell.text = str(column)
        for paragraph in cell.paragraphs:
            apply_paragraph_font(paragraph, size=7.0, bold=True)
    for _, row in preview.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = format_value(value)
            for paragraph in cells[idx].paragraphs:
                apply_paragraph_font(paragraph, size=7.0)
    doc.add_paragraph()


def write_reordered_tables() -> list[tuple[str, str, Path, int, int, int | None]]:
    outputs = []
    for number, title, builder, preview_rows in TABLE_SPECS:
        df = clean_dataframe(builder())
        path = write_csv(df, f"table_{number.lower()}_{slugify(title)}.csv")
        outputs.append((number, title, path, len(df), len(df.columns), preview_rows))
    return outputs


def slugify(text: str) -> str:
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")[:90]


def output_docx() -> Path:
    return Path(os.environ.get("LRI_SUPPLEMENT_OUTPUT_DOCX", str(DEFAULT_OUTPUT_DOCX)))


def build_docx(table_outputs: list[tuple[str, str, Path, int, int, int | None]]) -> Path:
    doc = Document()
    set_doc_style(doc)

    title = "Supplementary Appendix"
    p = add_para(doc, title, "Title", size=16, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = add_para(
        doc,
        "Reinterpreting high childhood lower respiratory infection burden through development-conditioned country profiles, 1990-2023",
        size=11,
        bold=True,
    )
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(doc, "Supplementary Methods", "Heading 1", size=13, bold=True)

    add_para(doc, "Supplementary Methods 1. Analytical scope and interpretation boundary", "Heading 2", size=12, bold=True)
    add_para(
        doc,
        "These supplementary materials support the main development-conditioned observed/excess typology analysis. They provide additional details on data sources, contextual-indicator coverage, expected-burden model validation, threshold and model sensitivity, age-boundary analyses, country-level membership stability, transition-pathway metrics, and contextual residual fingerprints. All analyses are ecological and location-level. The SDI-year expected burden is used as a model-based benchmark referent, not as a causal counterfactual or an estimate of avoidable burden. Contextual indicators are used for post-classification profiling and hypothesis generation, not for causal attribution. Exact location membership may vary across threshold rules, residual scales, and model specifications; interpretation therefore emphasises profile-level contrasts and transition pathways rather than deterministic country labels.",
    )

    add_para(doc, "Supplementary Methods 2. Expected-burden modelling and typology definition", "Heading 2", size=12, bold=True)
    add_para(
        doc,
        "Expected under-5 LRI incidence was estimated as an SDI-year benchmark using a spline term for SDI and calendar-year effects. Observed burden was classified using the upper tertile of observed rates in each analysis year. Development-adjusted excess burden was primarily defined on the additive excess-rate scale, with observed/expected and log observed/expected ratios retained for sensitivity and interpretation. We did not classify locations by posterior exceedance probabilities or prediction intervals because the local analysis package did not propagate GBD posterior draws through the benchmark model. WASH, UHC, vaccination, nutrition, HIV/TB, household energy, ambient PM2.5, and climate indicators were not included as primary benchmark-model covariates because the estimand was a development-conditioned referent rather than a multivariable causal adjustment model.",
    )

    add_para(doc, "Supplementary Methods 3. Contextual indicators and residual profiling", "Heading 2", size=12, bold=True)
    add_para(
        doc,
        "Contextual indicators were assembled from GBD, World Bank WDI, WHO GHO, ERA5, WorldPop, and related public data products. For each indicator, adverse-direction coding was defined so that higher residual z scores indicated a more adverse profile relative to same-year SDI-quintile peers. Variables were grouped into household energy, WASH, immunisation/UHC, HIV/TB, nutrition, health resources, ambient PM2.5, and climate/transmission-context domains. These residual fingerprints describe post-classification profiles; they are not ranked attributable contributors and should not be read as causal drivers.",
    )

    add_para(doc, "Supplementary Methods 4. Sensitivity and age-boundary analyses", "Heading 2", size=12, bold=True)
    add_para(
        doc,
        "Sensitivity analyses assessed threshold definitions, residual scales, model specifications, population-weighted WLS, contextual-profile contrast stability, membership stability across 84 under-5 scenarios, and age-boundary extensions. Additional robustness checks evaluated Duan smearing correction, 2000-anchored fixed thresholds, small-population influence using available under-5 population-at-risk, and harmonised 2018-2023 contextual windows. Under-5 LRI remained the primary outcome. All-age, 5-69 years, and 70+ analyses were used to test boundary behaviour, while granular 5-14, 15-49, and 50-69 analyses were retained as supplementary checks rather than as a second primary story.",
    )

    add_para(doc, "Supplementary Methods 5. Contextual-indicator grouping and adverse-direction coding", "Heading 2", size=12, bold=True)
    add_para(
        doc,
        "Lower access to clean cooking, basic water, basic sanitation, vaccination coverage, UHC, health expenditure, physician density, and nurse/midwife density was coded as adverse. Higher HAP-PM, ambient PM2.5, stunting, wasting, HIV prevalence, TB incidence, annual mean temperature, absolute humidity, and high-humidity days were coded as adverse. Outcome-side measures such as under-5 mortality were excluded from contextual residual profiling. Climate variables are interpreted as transmission-context descriptors, not acute weather effects.",
    )

    add_para(doc, "Supplementary Figures", "Heading 1", size=13, bold=True)
    doc.add_picture(str(SUPP_FIG_DIR / "supplementary_figure_s1_spatial_typology_2023.png"), width=Inches(8.8))
    add_para(
        doc,
        "Supplementary Figure S1. Spatial distribution of observed/excess typology and development-adjusted excess, 2023. The map is descriptive and should not be interpreted as a causal attributable-burden map.",
    )
    doc.add_picture(str(SUPP_FIG_DIR / "supplementary_figure_s2_annual_typology_counts.png"), width=Inches(8.2))
    add_para(
        doc,
        "Supplementary Figure S2. Annual observed/excess typology counts, 1990-2023. The figure shows year-specific relative typology composition using annual upper-tertile observed-incidence and excess-burden thresholds.",
    )
    doc.add_picture(str(SUPP_FIG_DIR / "supplementary_figure_s3_age_overlap_heatmap.png"), width=Inches(6.3))
    add_para(
        doc,
        "Supplementary Figure S3. Age-boundary high-observed/high-excess membership overlap. The figure summarises overlap in high-observed/high-excess membership across the primary under-5 outcome and older age-boundary analyses.",
    )

    add_para(doc, "Supplementary Tables", "Heading 1", size=13, bold=True)

    for number, title, path, rows, cols, preview_rows in table_outputs:
        add_para(doc, f"Supplementary Table {number}. {title}", "Heading 2", size=11, bold=True)
        df = read_csv(path)
        add_table_preview(doc, df, preview_rows)

    output = output_docx()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    with zipfile.ZipFile(output) as package:
        bad = package.testzip()
        if bad:
            raise RuntimeError(f"Invalid DOCX package member: {bad}")
    return output


def main() -> None:
    from build_robustness_supplement_s19_s22 import main as build_robustness_tables
    from build_supplementary_figures import main as build_supplementary_figures

    build_robustness_tables()
    table_outputs = write_reordered_tables()
    build_supplementary_figures()
    output = build_docx(table_outputs)
    print(f"Wrote {output}")
    print(f"Wrote reordered CSV tables to {OUTPUT_TABLE_DIR}")


if __name__ == "__main__":
    main()
